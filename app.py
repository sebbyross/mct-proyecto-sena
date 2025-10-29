from flask import Flask, render_template, request, redirect, url_for, session, flash, make_response, jsonify, send_file
import pymysql
from io import BytesIO
from reportlab.pdfgen import canvas
from flask_bcrypt import Bcrypt
from openpyxl import Workbook
import os
from werkzeug.utils import secure_filename 
from werkzeug.security import generate_password_hash

app = Flask(__name__)
app.secret_key = 'tu_clave_secreta'

# SQL CONFIG---------------------------------------------------------
app.config['MYSQL_HOST'] = os.getenv('MYSQL_HOST', 'localhost')
app.config['MYSQL_USER'] = os.getenv('MYSQL_USER', 'root')
app.config['MYSQL_PASSWORD'] = os.getenv('MYSQL_PASSWORD', '')
app.config['MYSQL_DB'] = os.getenv('MYSQL_DB', 'flask_login')

def get_db_connection():
    return pymysql.connect(
        host=app.config['MYSQL_HOST'],
        user=app.config['MYSQL_USER'],
        password=app.config['MYSQL_PASSWORD'],
        database=app.config['MYSQL_DB'],)
bcrypt = Bcrypt(app)
# UTIL CUSTOM FUNCTIONS---------------------------------------------------------

def add_no_cache_headers(response):
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache" 
    response.headers["Expires"] = "0"
    return response

def list_whole_db():
    products_count = users_count = low_stock_count = 0
    productos = []
    users = []
    recent_products = []
    conn = get_db_connection()
    cur = conn.cursor()

    cur.execute('SELECT COUNT(*) FROM productos')
    r = cur.fetchone()
    products_count = int(r[0]) if r and r[0] is not None else 0

    cur.execute('SELECT COUNT(*) FROM users')
    r = cur.fetchone()
    users_count = int(r[0]) if r and r[0] is not None else 0

    cur.execute('SELECT id, nombre, descripcion, precio, stock FROM productos ORDER BY id DESC LIMIT 5')
    recent_products = cur.fetchall() or []

    cur.execute('SELECT id, nombre, descripcion, precio, stock FROM productos')
    productos = cur.fetchall() or []

    # low-stock threshold (adjust number if you want)
    cur.execute('SELECT COUNT(*) FROM productos WHERE stock < %s', (20,))
    r = cur.fetchone()
    low_stock_count = int(r[0]) if r and r[0] is not None else 0

    cur.execute('SELECT id, username, password, adminstatus FROM users')
    users = cur.fetchall() or []
    return {
        "products_count": products_count,
        "users_count": users_count,
        "low_stock_count": low_stock_count,
        "productos": productos,
        "users": users,
        "recent_products": recent_products
    }

# RUTAS---------------------------------------------------------

@app.route('/')
def home():
    if 'username' in session:
        return redirect(url_for('index'))
    return redirect(url_for('login')) 

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']

        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute('SELECT * FROM users WHERE username = %s', (username,))
        user = cur.fetchone()
        id = user[0] if user else None  # Obtener el ID del usuario si existe
        cur.close()

        if user and bcrypt.check_password_hash(user[2], password):
            session['username'] = username
            session['id'] = id  # Guardar el ID del usuario en la sesión

            conn = get_db_connection()
            cur = conn.cursor()
            cur.execute("SELECT adminstatus FROM users WHERE username = %s", (username,))
            adminstatus_result = cur.fetchone()
            cur.close()
            session['adminstatus'] = int(adminstatus_result[0]) if adminstatus_result and adminstatus_result[0] is not None else 0

            # Chequear rol
            if adminstatus_result and adminstatus_result[0] == 1:
                return redirect(url_for('admin'))
            else:
                return redirect(url_for('index'))
        else:
            flash('Usuario o contraseña incorrectos')
            return redirect(url_for('login'))

    response = make_response(render_template('login.html'))
    return add_no_cache_headers(response)

@app.route('/regis', methods=['GET', 'POST'])
def regis():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        adminstatus = 1 if request.form.get('adminstatus', '0') == '1' else 0

        if not username or not password:
            flash('Completa todos los campos')
            return redirect(url_for('regis'))

        # hash password (prefer bcrypt if configured)
        if 'bcrypt' in globals():
            hashed = bcrypt.generate_password_hash(password).decode('utf-8')
        else:
            hashed = generate_password_hash(password)

        try:
            conn = get_db_connection()
            cur = conn.cursor()
            cur.execute('INSERT INTO users (username, password, adminstatus) VALUES (%s, %s, %s)',
                        (username, hashed, adminstatus))
            conn.commit()
            cur.close()
            conn.close()
            flash('Usuario registrado correctamente')
            # if admin created user, go to users list, otherwise to login
            if session.get('adminstatus') == 1:
                return redirect(url_for('user_list'))
            return redirect(url_for('login'))
        except Exception as e:
            # keep message short; you can log e elsewhere
            flash('Error al crear el usuario')
            return redirect(url_for('regis'))

    return render_template('regis.html')
@app.route('/index')
def index():
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT adminstatus FROM users WHERE username = %s", (session['username'],))
    adminstatus_result = cur.fetchone()
    cur.close()
    if adminstatus_result and adminstatus_result[0] == 1:
        return redirect(url_for('admin'))

    data = list_whole_db()

    if 'username' not in session:
        return redirect(url_for('login'))
    response = make_response(render_template('index.html', **data))
    return add_no_cache_headers(response)

@app.route('/productos_user')
def productos_user():

    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT * FROM productos")
    productos = cur.fetchall()
    cur.close()
    return render_template('productos-user.html', productos=productos)

# FUNCIONES DE ADMINISTRADOR ------------------------------------------------------

@app.route('/admin')
def admin():
    # require admin
    if session.get('adminstatus') != 1:
        return redirect(url_for('index'))

    data = list_whole_db()

    return render_template('admin.html', **data)

@app.route('/users')
def user_list():
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT * FROM users")
    users = cur.fetchall()
    cur.close()

    return render_template('users.html', users=users)

@app.route('/search', methods=['GET', 'POST'])
def search_by_id():
    if request.method == 'POST':
        user_id = request.form['id']  # Get the ID from the form input

        # Query the database for the record with the given ID
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("SELECT * FROM users WHERE id = %s", (user_id,))
        user = cur.fetchone()
        cur.close()

        if user:
            # If a user is found, display the result
            return render_template('search_result.html', user=user)
        else:
            # If no user is found, flash a message
            flash("No hay usuario con esta id.", "warning")
            return redirect(url_for('search_by_id'))

    return render_template('search-user.html')

@app.route('/search-p', methods=['GET', 'POST'])
def search_p_by_id():
    if request.method == 'POST':
        product_id = request.form['id']  # Get the ID from the form input

        # Query the database for the record with the given ID
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("SELECT * FROM productos WHERE id = %s", (product_id,))
        producto = cur.fetchone()
        cur.close()

        if producto:
            # If a product is found, display the result
            return render_template('search_result_p.html', producto=producto)
        else:
            # If no user is found, flash a message
            flash("No se encontro el producto", "warning")
            return redirect(url_for('search_p_by_id'))

    return render_template('search-p.html')

@app.route('/productos')
def productos():
    
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT * FROM productos")
    productos = cur.fetchall()
    cur.close()

    return render_template('productos.html', productos=productos)

@app.route('/productos/agregar', methods=['GET', 'POST'])
def agregar_producto():

    if request.method == 'POST':
        nombre = request.form['nombre']
        descripcion = request.form['descripcion']
        precio = float(request.form['precio'])
        stock = int(request.form['stock'])

        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO productos (nombre, descripcion, precio, stock) VALUES (%s, %s, %s, %s)",
            (nombre, descripcion, precio, stock)
        )
        conn.commit()
        cur.close()

        return redirect(url_for('productos'))

    return render_template('productos-agregar.html')


@app.route('/productos/editar/<int:id>', methods=['GET', 'POST'])
def editar_producto(id):

    conn = get_db_connection()
    cur = conn.cursor()
    if request.method == 'POST':
        nombre = request.form['nombre']
        descripcion = request.form['descripcion']
        precio = request.form['precio']
        stock = request.form['stock']

        cur.execute("""
            UPDATE productos SET nombre=%s, descripcion=%s, precio=%s, stock=%s WHERE id=%s
        """, (nombre, descripcion, precio, stock, id))
        conn.commit()
        cur.close()
        return redirect(url_for('productos'))

    cur.execute("SELECT * FROM productos WHERE id=%s", (id,))
    producto = cur.fetchone()
    cur.close()

    return render_template('productos-editar.html', producto=producto)


@app.route('/productos/eliminar/<int:id>')
def eliminar_producto(id):

    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("DELETE FROM productos WHERE id=%s", (id,))
    conn.commit()
    cur.close()

    return redirect(url_for('productos'))



@app.route('/contact')
def contact():
    if 'username' not in session:
        return redirect(url_for('login'))
    response = make_response(render_template('contact.html'))
    return add_no_cache_headers(response)

@app.route('/more')
def more():
    if 'username' not in session:
        return redirect(url_for('login'))
    response = make_response(render_template('more.html'))
    return add_no_cache_headers(response)

@app.route('/logout')
def logout():
    session.pop('username', None)
    response = make_response(redirect(url_for('login')))
    return add_no_cache_headers(response)

@app.route('/users/editar/<int:id>', methods=['GET', 'POST'])
def editar_usuario(id):
    conn = get_db_connection()
    cur = conn.cursor()
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()

        # safe retrieval: checkbox may be absent when unchecked
        adminstatus_raw = request.form.get('adminstatus', '0')
        adminstatus = 1 if str(adminstatus_raw) == '1' else 0

        # optional: prevent an admin from removing their own admin flag
        if session.get('id') and int(session.get('id')) == int(id) and adminstatus == 0:
            flash('No puedes quitarte permisos de administrador a ti mismo.')
            return redirect(url_for('user_list'))

        if password:
            hashed_password = bcrypt.generate_password_hash(password).decode('utf-8')
            cur.execute("""
                UPDATE users SET username=%s, password=%s, adminstatus=%s WHERE id=%s
            """, (username, hashed_password, adminstatus, id))
        else:
            cur.execute("""
                UPDATE users SET username=%s, adminstatus=%s WHERE id=%s
            """, (username, adminstatus, id))

        conn.commit()
        cur.close()
        return redirect(url_for('user_list'))

    cur.execute("SELECT * FROM users WHERE id=%s", (id,))
    user = cur.fetchone()
    cur.close()

    return render_template('users-editar.html', user=user)


@app.route('/users/eliminar/<int:id>')
def eliminar_usuario(id):
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("DELETE FROM users WHERE id=%s", (id,))
    conn.commit()
    cur.close()
    conn.close() 
    return redirect(url_for('user_list'))

# REPORTES ------------------------------------------------------

from openpyxl.drawing.image import Image as XLImage
import tempfile
import os

@app.route('/report/excel')
def report_excel():
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT * FROM productos")
    data = cur.fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(["Id", "Nombre", "Descripción", "Precio", "Stock"])

    row_num = 2
    for id, nombre, descripcion, precio, stock in data:
        ws.cell(row=row_num, column=1, value=id)
        ws.cell(row=row_num, column=2, value=nombre)
        ws.cell(row=row_num, column=3, value=descripcion)
        ws.cell(row=row_num, column=4, value=precio)
        ws.cell(row=row_num, column=5, value=stock)
        row_num += 1

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(output, download_name="reporte_productos.xlsx", as_attachment=True)

@app.route('/report/pdf')
def report_pdf():
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT * FROM productos")
    data = cur.fetchall()

    buffer = BytesIO()
    p = canvas.Canvas(buffer, pagesize=(595, 842))  # A4
    y = 800

    p.setFont("Helvetica-Bold", 16)
    p.drawString(200, y, "Reporte de Productos")
    y -= 40

    for id, nombre, descripcion, precio, stock in data:
        if y < 100:
            p.showPage()
            y = 800

        p.setFont("Helvetica-Bold", 12)
        p.drawString(50, y, f"ID: {id}")
        y -= 15
        p.drawString(50, y, f"Nombre: {nombre}")
        y -= 15
        p.setFont("Helvetica", 10)
        p.drawString(50, y, f"Descripción: {descripcion}")
        y -= 15
        p.drawString(50, y, f"Precio: {precio}")
        y -= 15
        p.drawString(50, y, f"Stock: {stock}")
        y -= 15

    p.save()
    buffer.seek(0)
    return send_file(buffer, download_name="reporte_productos.pdf", as_attachment=True)

@app.route('/account', methods=['GET', 'POST'])
def account():
    if 'username' not in session:
        return redirect(url_for('login'))

    message = None
    if request.method == 'POST':
        new_username = request.form.get('username', '').strip()
        new_password = request.form.get('password', '').strip()

        # Try to persist changes using your get_db_connection() helper.
        try:
            conn = get_db_connection()
            cur = conn.cursor()
            if new_password:
                # prefer flask-bcrypt when available
                if 'bcrypt' in globals():
                    hashed = bcrypt.generate_password_hash(new_password).decode('utf-8')
                else:
                    hashed = generate_password_hash(new_password)
                cur.execute('UPDATE users SET username=%s, password=%s WHERE id=%s',
                            (new_username or session.get('username'), hashed, session.get('id')))
            else:
                cur.execute('UPDATE users SET username=%s WHERE id=%s',
                            (new_username or session.get('username'), session.get('id')))
            conn.commit()
            cur.close()
            conn.close()
            message = 'Datos actualizados.'
            if new_username:
                session['username'] = new_username
        except Exception:
            if new_username:
                session['username'] = new_username
                message = 'Usuario actualizado en sesión (sin persistencia).'
            else:
                message = 'Sin cambios (no hay BD disponible).'

    # create user object for template
    user = type('U', (), {})()
    user.id = session.get('id')
    user.username = session.get('username')
    user.adminstatus = session.get('adminstatus', 0)

    return render_template('account.html', user=user, message=message)

if __name__ == '__main__':
    app.run(debug=True)